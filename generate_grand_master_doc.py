import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_sec_header(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79) # Deep Navy
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
    elif level == 2:
        run.font.size = Pt(13.5)
        run.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32) # Forest Green
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
    elif level == 3:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x6A, 0x1B, 0x9A) # Purple
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
    return p

def add_justified_para(doc, text, bold_prefix="", highlight_words=[], space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.size = Pt(12)
        r_pre.font.name = 'Calibri'
        r_pre.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        
    r_text = p.add_run(text)
    r_text.font.size = Pt(12)
    r_text.font.name = 'Calibri'
    r_text.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return p

def make_callout(doc, title, text, bg_hex="F0F4F8", border_hex="1F4E79"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=160, bottom=160, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(4)
    
    run_t = p.add_run(f"{title}\n")
    run_t.bold = True
    run_t.font.size = Pt(12)
    run_t.font.name = 'Calibri'
    run_t.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    
    run_txt = p.add_run(text)
    run_txt.font.size = Pt(11)
    run_txt.font.name = 'Calibri'
    run_txt.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(6)

def build_grand_master_doc():
    doc = docx.Document()
    
    # 0.75 inch margins for full edge-to-edge layout & high text capacity per page
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Calibri'
    normal_font.size = Pt(12)
    normal_font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    # Title
    t_p = doc.add_paragraph()
    t_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t_run = t_p.add_run("ISML PORTAL & LMS PLATFORM — GRAND MASTER PRODUCT DEVELOPMENT & DELIVERY ROADMAP")
    t_run.bold = True
    t_run.font.size = Pt(20)
    t_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    t_p.paragraph_format.space_after = Pt(4)
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Complete 3-Product Architecture, Technical Specifications, Responsibility Division, & Release Schedule")
    sub_run.font.size = Pt(13)
    sub_run.italic = True
    sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    sub_p.paragraph_format.space_after = Pt(16)
    
    make_callout(
        doc,
        "EXECUTIVE TIMELINE & DIWALI DEADLINE COMMITMENT REQUEST",
        "This document establishes the master delivery plan for Product 1 (ISML Portal V3 - Targeted for October 4, 2026), "
        "Product 2 (College LMS Edition - Targeted for October 31, 2026), and Product 3 (School LMS Edition - Targeted for October 31, 2026). "
        "To safeguard product stability against unforeseen production roadblocks, a strict 1-Week Contingency Backup Window is included. "
        "The team formally commits to handing over all three complete, production-verified products prior to the upcoming Diwali Holidays.",
        bg_hex="FFF8E1",
        border_hex="FF8F00"
    )

    # -------------------------------------------------------------
    # SECTION 1: EXECUTIVE ROADMAP & MILESTONE SCHEDULE
    # -------------------------------------------------------------
    add_sec_header(doc, "1. Executive Master Schedule & Release Timelines", level=1)
    
    add_justified_para(
        doc,
        "The ISML engineering ecosystem encompasses three major commercial product lines: Product 1 (ISML Portal V3), Product 2 (College LMS Edition), and Product 3 (School LMS Edition). In order to maintain strict engineering quality while fulfilling operational deployment demands, a phased milestone schedule has been established.",
        bold_prefix="Multi-Product Delivery Strategy: "
    )
    
    # Milestone Schedule Table
    t_sched = doc.add_table(rows=1, cols=4)
    t_sched.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sched.autofit = False
    set_table_borders(t_sched)
    
    s_hdr = t_sched.rows[0].cells
    s_widths = [Inches(1.8), Inches(1.8), Inches(1.8), Inches(1.6)]
    for i, title in enumerate(["Product Name / Edition", "Target Completion Date", "Contingency / Backup Buffer", "Diwali Handover Status"]):
        s_hdr[i].width = s_widths[i]
        set_cell_background(s_hdr[i], "1F4E79")
        set_cell_margins(s_hdr[i])
        p = s_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    s_data = [
        ("Product 1: ISML Portal V3", "October 4, 2026 (Sunday)", "1-Week Buffer (October 11)", "Guaranteed pre-Diwali Handover"),
        ("Product 2: College LMS Edition", "October 31, 2026 (Saturday)", "1-Week Buffer (November 7)", "Guaranteed pre-Diwali Handover"),
        ("Product 3: School LMS Edition", "October 31, 2026 (Saturday)", "1-Week Buffer (November 7)", "Guaranteed pre-Diwali Handover")
    ]
    for p_name, t_date, c_buf, d_stat in s_data:
        r_cells = t_sched.add_row().cells
        for i, val in enumerate([p_name, t_date, c_buf, d_stat]):
            r_cells[i].width = s_widths[i]
            set_cell_background(r_cells[i], "F0F4F8")
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # SECTION 2: HIGH-LEVEL TEAM RESPONSIBILITY DIVISION
    # -------------------------------------------------------------
    add_sec_header(doc, "2. High-Level Team Ownership & Resource Division", level=1)
    
    add_justified_para(
        doc,
        "The technical execution of the entire 3-product portfolio is organized around a Lead Technical Owner model combined with a dedicated Student & Quality Assurance Owner. Jay Kumar serves as the Lead Development Owner managing architecture, full-stack implementation, database schema, APIs, live class engines, and DevOps across all 3 products. Bharathi serves as the Primary Developer for the Student Module and Lead QA / Testing Auditor responsible for testing all 3 products end-to-end.",
        bold_prefix="Master Team Model: "
    )
    
    t_team = doc.add_table(rows=1, cols=4)
    t_team.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_team.autofit = False
    set_table_borders(t_team)
    
    tm_hdr = t_team.rows[0].cells
    tm_widths = [Inches(1.5), Inches(1.8), Inches(1.2), Inches(2.5)]
    for i, title in enumerate(["Team Member", "Primary Development Role", "QA Audit Role", "Assigned Product Responsibilities"]):
        tm_hdr[i].width = tm_widths[i]
        set_cell_background(tm_hdr[i], "1F4E79")
        set_cell_margins(tm_hdr[i])
        p = tm_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    tm_data = [
        ("Jay Kumar", "Lead Technical Architect & Full-Stack Owner", "Self Verification & Code Review", "Product 1 (Live Class, Recording, Schedule, Video Library), Product 2 (College Core, Super Admin, Manager, Finance, Faculty), Product 3 (School Core, Grades, Timetable, Career AI)."),
        ("Bharathi", "Resource Agent Owner & Student Developer", "Lead QA & Testing Auditor", "Product 1 (Resource Agent Intelligence Engine, LSRW Speech Agent), Product 2 & 3 (Student Module UI Development, End-to-End QA Testing across all 3 products)."),
        ("Interns + Bharathi", "LSRW Activity & Content Support", "Activity Testing", "Assisting Bharathi in implementing Listening, Speaking, Reading, and Writing practice datasets, audio parsing, and AI prompt validation.")
    ]
    for name, dev_role, qa_role, scope in tm_data:
        r_cells = t_team.add_row().cells
        bg_col = "F0F4F8" if "Jay" in name else ("E8F5E9" if "Bharathi" == name else "F3E5F5")
        txt_c = RGBColor(0x1F, 0x4E, 0x79) if "Jay" in name else (RGBColor(0x2E, 0x7D, 0x32) if "Bharathi" == name else RGBColor(0x6A, 0x1B, 0x9A))
        for i, val in enumerate([name, dev_role, qa_role, scope]):
            r_cells[i].width = tm_widths[i]
            set_cell_background(r_cells[i], bg_col)
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = txt_c
                
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # SECTION 3: PRODUCT 1 — ISML PORTAL V3 DEEP DIVE
    # -------------------------------------------------------------
    add_sec_header(doc, "3. Product 1: ISML Portal V3 — Comprehensive Technical Specifications", level=1)
    
    add_justified_para(
        doc,
        "Product 1 (ISML Portal V3) represents the interactive real-time core of the ecosystem. It encompasses three massive functional pillars: Jay Kumar's Live Class & Recording System, Bharathi's AI Resource Agent, and the collaborative Bharathi + Interns LSRW Speech & Language Engine. Target Completion Date: October 4, 2026.",
        bold_prefix="Product 1 Focus: "
    )
    
    # Pillar 3.1: Jay Kumar Live System
    add_sec_header(doc, "3.1 Jay Kumar — Live Learning & Recording Ecosystem", level=2)
    add_justified_para(
        doc,
        "Jay Kumar owns the complete lifecycle of live class operations, WebRTC video streaming, interactive classroom controls, attendance logging, video encoding, and pre-recorded library delivery.",
        bold_prefix="Scope Overview: "
    )
    
    p1_jay_table = doc.add_table(rows=1, cols=3)
    p1_jay_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    p1_jay_table.autofit = False
    set_table_borders(p1_jay_table)
    
    pj_hdr = p1_jay_table.rows[0].cells
    pj_widths = [Inches(2.0), Inches(2.0), Inches(3.0)]
    for i, title in enumerate(["Module Component", "Functional Focus", "Detailed Technical Implementation"]):
        pj_hdr[i].width = pj_widths[i]
        set_cell_background(pj_hdr[i], "1F4E79")
        set_cell_margins(pj_hdr[i])
        p = pj_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    p1_j_rows = [
        ("1. Class Scheduling & Lifecycle", "Session State Engine", "Configuring course, batch, teacher, date, time, and duration. Managing states: Draft -> Scheduled -> Ready -> Live -> Completed -> Recording Processing -> Recording Available. Handling failure states (Cancelled, Failed)."),
        ("2. Live Room Security & Access", "Token & Security Service", "Generating unique session/room IDs, issuing dynamic Daily.co/Agora JWT access tokens, enforcing role authentication (Host vs Student), and handling network reconnections."),
        ("3. Teacher Live Control Room", "Interactive Host Interface", "Building classroom controls: Start/End class, Mute All, Unmute Student, Screen Sharing, Live Text Chat, Raise Hand alerts, Participant management, and Recording status indicators."),
        ("4. Student Live Classroom", "Learner Stream UI", "Upcoming class view, one-click secure join, waiting room validation, teacher stream video player, in-class chat participation, and reconnection handling."),
        ("5. Attendance Automation Engine", "Activity-Based Logging", "Automated join/leave timestamp capture, disconnect tracking, participation duration aggregation, and automatic Present/Absent calculation without manual teacher input."),
        ("6. Session Monitoring Dashboard", "Real-Time Operational View", "Management dashboard tracking active live classes, participant counts, teacher connection health, session duration metrics, and recording pipeline status."),
        ("7. Recording & Processing Pipeline", "Cloud Encoding Engine", "Automated recording generation via cloud webhooks, availability polling, metadata extraction, retry handling upon processing failure, and validation."),
        ("8. Recording to Pre-recorded Sync", "Zero-Touch Automation", "Automated workflow: Live Class Completed -> Recording Generated -> Read Metadata (Course, Batch, Topic, Teacher) -> Create Pre-Recorded Video item in Library."),
        ("9. Pre-recorded Video Library", "Secure Course Repository", "Course, batch, and topic indexing, secure tokenized video playback URLs, student watch progress tracking, resume position, and completion percentage calculation.")
    ]
    for comp, foc, impl in p1_j_rows:
        r_cells = p1_jay_table.add_row().cells
        for i, val in enumerate([comp, foc, impl]):
            r_cells[i].width = pj_widths[i]
            set_cell_background(r_cells[i], "F0F4F8")
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Pillar 3.2: Bharathi Resource Agent
    add_sec_header(doc, "3.2 Bharathi — AI Resource Agent & Intelligence Engine", level=2)
    add_justified_para(
        doc,
        "Bharathi owns the complete Resource Agent ecosystem. This module executes automated web discovery, academic classification, deduplication, copyright risk analysis, and human review queue preparation.",
        bold_prefix="Resource Intelligence Scope: "
    )
    
    p1_b_table = doc.add_table(rows=1, cols=3)
    p1_b_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    p1_b_table.autofit = False
    set_table_borders(p1_b_table)
    
    pb_hdr = p1_b_table.rows[0].cells
    pb_widths = [Inches(2.0), Inches(2.0), Inches(3.0)]
    for i, title in enumerate(["Intelligence Stage", "Processing Component", "Detailed Technical Execution"]):
        pb_hdr[i].width = pb_widths[i]
        set_cell_background(pb_hdr[i], "2E7D32")
        set_cell_margins(pb_hdr[i])
        p = pb_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    p1_b_rows = [
        ("1. Academic Requirement Parser", "Curriculum Taxonomy Mapping", "Parsing target Language, Course, Level, Category, Skill, Topic, Learning Objective, required Resource Type, and Academic Context."),
        ("2. Resource Web Discovery", "External Web Search Crawler", "Dynamic search query generation, web discovery, candidate filtering, URL harvesting, and source provider verification."),
        ("3. Metadata Extraction", "Automated Scraping Pipeline", "Extracting title, description, source, provider, language, level, topic, skill, resource type, and URL structure."),
        ("4. Academic Classification", "Hierarchy Mapping Engine", "Structuring resources into Language -> Course -> Level -> Category -> Skill -> Topic -> Resource Type taxonomy."),
        ("5. URL & Health Validation", "Link Integrity Checker", "Detecting broken URLs, HTTP redirects, page accessibility, content availability, and metadata completeness."),
        ("6. Multi-Stage Deduplication", "Deterministic Matcher", "Checking Exact URL -> Canonical URL -> Normalized URL -> URL Hash -> Content Similarity scoring."),
        ("7. Quality & Relevance Analysis", "Evidence-Based Scoring", "Evaluating resource quality and academic relevance. Strict Rule: No AI guessing without empirical evidence; flags uncertainty for human audit."),
        ("8. Copyright Risk Analysis", "Compliance Sentinel", "Analyzing license signals, attribution requirements, and usage restrictions. Routes ambiguous cases to Human Review Queue."),
        ("9. Human Review Queue", "Editorial Approval Interface", "Preparing AI analysis reports for Human Review (Approve / Modify / Reject) prior to publishing into main library."),
        ("10. Bulk Import & Reconciliation", "Batch Import Processor", "Importing tutor-contributed resources, executing row-by-row deduplication reconciliation, and logging failed import rows.")
    ]
    for comp, foc, impl in p1_b_rows:
        r_cells = p1_b_table.add_row().cells
        for i, val in enumerate([comp, foc, impl]):
            r_cells[i].width = pb_widths[i]
            set_cell_background(r_cells[i], "E8F5E9")
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Pillar 3.3: Bharathi + Interns LSRW
    add_sec_header(doc, "3.3 Bharathi + Interns — LSRW Speech & Language Agent", level=2)
    add_justified_para(
        doc,
        "Bharathi and the intern team co-own the LSRW Speech & Language Engine. This module delivers interactive Listening, Speaking, Reading, and Writing practice activities powered by AI evaluation.",
        bold_prefix="LSRW Ecosystem Scope: "
    )
    
    p1_l_table = doc.add_table(rows=1, cols=3)
    p1_l_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    p1_l_table.autofit = False
    set_table_borders(p1_l_table)
    
    pl_hdr = p1_l_table.rows[0].cells
    pl_widths = [Inches(2.0), Inches(2.0), Inches(3.0)]
    for i, title in enumerate(["LSRW Pillar", "Module Focus", "Technical Implementation & AI Pipeline"]):
        pl_hdr[i].width = pl_widths[i]
        set_cell_background(pl_hdr[i], "6A1B9A")
        set_cell_margins(pl_hdr[i])
        p = pl_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    p1_l_rows = [
        ("1. Listening Module", "Audio Comprehension", "Level/Topic-based audio playback generation, comprehension question creation, student attempt logging, AI evaluation of answers, and feedback generation."),
        ("2. Speaking Agent", "Voice AI & Phonetics Engine", "Student audio recording capture, Speech-to-Text parsing. AI Evaluation: Pronunciation, Fluency, Grammar, Vocabulary, Sentence Structure, and Relevance scoring."),
        ("3. Reading Module", "Text Passages & Analysis", "Level-appropriate reading content generation, vocabulary-focused passages, comprehension questions, reading score calculation, and feedback."),
        ("4. Writing Module", "Essay & Sentence Analysis", "Writing prompt generation, AI language analysis (Grammar, Vocabulary, Coherence, Sentence Structure, Relevance), instant scoring, and detailed feedback."),
        ("5. Unified Dashboard", "Progress Tracking", "Aggregating Listening, Speaking, Reading, and Writing scores into a unified AI Academic Progress Dashboard.")
    ]
    for comp, foc, impl in p1_l_rows:
        r_cells = p1_l_table.add_row().cells
        for i, val in enumerate([comp, foc, impl]):
            r_cells[i].width = pl_widths[i]
            set_cell_background(r_cells[i], "F3E5F5")
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x6A, 0x1B, 0x9A)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # SECTION 4: PRODUCT 2 (COLLEGE LMS) & PRODUCT 3 (SCHOOL LMS) DEEP DIVE
    # -------------------------------------------------------------
    add_sec_header(doc, "4. Product 2 (College LMS) & Product 3 (School LMS) Specifications", level=1)
    
    add_justified_para(
        doc,
        "Product 2 (College LMS Edition) and Product 3 (School LMS Edition) share approximately 90% of their core platform infrastructure while maintaining specialized academic structures. Target Completion Date for both editions: October 31, 2026.",
        bold_prefix="Dual LMS Overview: "
    )
    
    # Comparative Product Table
    t_lms_comp = doc.add_table(rows=1, cols=3)
    t_lms_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_lms_comp.autofit = False
    set_table_borders(t_lms_comp)
    
    lc_hdr = t_lms_comp.rows[0].cells
    lc_widths = [Inches(2.0), Inches(2.5), Inches(2.5)]
    for i, title in enumerate(["Functional Module", "Product 2: College LMS Edition", "Product 3: School LMS Edition"]):
        lc_hdr[i].width = lc_widths[i]
        set_cell_background(lc_hdr[i], "1F4E79")
        set_cell_margins(lc_hdr[i])
        p = lc_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    lms_comp_data = [
        ("Academic Hierarchy", "Department -> Program -> Semester -> Batch -> Subject", "Grade / Standard -> Section -> Subject -> Class"),
        ("Scheduling Engine", "Semester time-slot allocation, Faculty credit hours", "Period-based daily timetable, Class Teacher assignment"),
        ("User Roles", "Dean, HOD, Professor, Assistant Professor, Student", "Principal, Vice-Principal, Class Teacher, Subject Teacher, Student"),
        ("Attendance System", "Subject-wise & Semester-wise percentage tracking", "Daily morning/afternoon session attendance & Subject attendance"),
        ("Special AI Modules", "Higher Education Research AI & Skill Mapping", "K-12 Career Guidance AI & Interactive Speaking Bot"),
        ("Development Owner", "Jay Kumar (Shared Core & College Workflow)", "Jay Kumar (Shared Core & School Workflow)"),
        ("Student & QA Owner", "Bharathi (Student UI & College QA Audit)", "Bharathi (Student UI & School QA Audit)")
    ]
    for mod, col_spec, sch_spec in lms_comp_data:
        r_cells = t_lms_comp.add_row().cells
        for i, val in enumerate([mod, col_spec, sch_spec]):
            r_cells[i].width = lc_widths[i]
            set_cell_background(r_cells[i], "F0F4F8")
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Detailed LMS Roles & Modules Matrix
    add_sec_header(doc, "4.1 Detailed LMS Role & Sub-system Execution Matrix", level=2)
    
    lms_role_matrix = [
        ("1. Super Admin Module", "Jay Kumar (Development)", "Platform dashboard, multi-institution onboarding (College vs School), global roles, system audit logs, and infrastructure monitoring.", "Bharathi (QA Testing)", "Validating institution creation, privilege escalation security, audit log visibility, and global configuration overrides."),
        ("2. Manager Module", "Jay Kumar (Development)", "User account creation, User ID generation, teacher/faculty management, class monitoring dashboards, and operational approvals.", "Bharathi (QA Testing)", "Testing user creation workflows, role assignment, class monitoring filters, operational approval triggers, and access limits."),
        ("3. Finance Manager", "Jay Kumar (Development)", "Fee structure definition, student fee mapping, payment status tracking, receipt PDF generation, Razorpay/Stripe webhook integration.", "Bharathi (QA Testing)", "Testing fee creation, payment gateway callbacks, webhook processing, receipt generation, and transaction reconciliation."),
        ("4. Academic Manager", "Jay Kumar (Development)", "Hierarchy management (Language -> Course -> Level -> Topic -> Skill), resource taxonomy mapping, and curriculum publishing.", "Bharathi (QA Testing)", "Testing academic taxonomy integrity, course level mapping, skill tagging, and resource publishing workflows."),
        ("5. Teacher / Faculty Module", "Jay Kumar (Development)", "Teacher dashboard, assigned batches, class scheduling, live classroom controls, assignment creation, exam management, and marks entry.", "Bharathi (QA Testing)", "Testing teacher login, batch filtering, live room host controls, assignment publishing, marks entry, and student feedback."),
        ("6. Assistant Teacher Module", "Jay Kumar (Development)", "Assigned class monitoring, student attendance support, assignment review assistance, and classroom moderation.", "Bharathi (QA Testing)", "Testing assistant permission boundaries (cannot alter primary grades), attendance logging verification, and teacher support."),
        ("7. Doubt & Assistant Doubt", "Jay Kumar (Development)", "Real-time doubt queue, course/topic categorization, student question inspection, text/voice answers, and escalation workflows.", "Bharathi (QA Testing)", "Testing doubt submission receipt, queue prioritization, escalation handoff, student notification delivery, and status updates."),
        ("8. Student Module", "Bharathi (Development Owner)", "Student dashboard UI, My Courses, live class join, pre-recorded video player, assignment submission, online exams, and profile.", "Jay Kumar (APIs) & Bharathi (QA)", "Jay delivers backend APIs & WebSockets. Bharathi builds client UI, tests responsive layouts, validates UX, and runs regression tests."),
        ("9. Common Scheduling Engine", "Jay Kumar (Development)", "Unified matrix handling College (Department/Semester) & School (Grade/Period) scheduling with conflict detection algorithms.", "Bharathi (QA Testing)", "Testing schedule creation, time slot conflict detection, teacher overlap prevention, student timetable views, and rescheduling."),
        ("10. Live Classes Engine", "Jay Kumar (Development)", "WebRTC / Daily.co / Agora API integration, secure room tokens, teacher/student authentication, real-time video grid, and duration logging.", "Bharathi (QA Testing)", "Testing multi-user room entry, token expiration security, camera/mic toggles, screen share quality, and network reconnection."),
        ("11. Attendance Automation", "Jay Kumar (Development)", "Activity-based automated tracking: Join/leave timestamp capture, disconnect aggregation, and automatic Present/Absent status updating.", "Bharathi (QA Testing)", "Testing join/leave timestamp accuracy, late-join threshold logic, early-leave detection, and attendance report correctness."),
        ("12. Recording & Pre-Recorded", "Jay Kumar (Development)", "Automated recording pipeline: Live Class Completed -> Cloud Processing -> Video Library Indexing -> Secure HLS Playback Player.", "Bharathi (QA Testing)", "Testing recording webhooks, metadata mapping accuracy, video library indexing, secure URL authorization, and resume position."),
        ("13. AI & Intelligent Engine", "Jay Kumar (Development)", "AI Doubt Solver, course-context RAG search, Speaking Practice Bot (voice evaluation), and K-12 Career Guidance AI (Product 3).", "Bharathi (QA Testing)", "Testing AI response latency, prompt accuracy, voice transcription quality in Speaking Bot, feedback rendering, and regression impact.")
    ]
    
    for r_title, d_owner, d_desc, q_owner, q_desc in lms_role_matrix:
        add_sec_header(doc, r_title, level=3)
        t_rm = doc.add_table(rows=1, cols=3)
        t_rm.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_rm.autofit = False
        set_table_borders(t_rm)
        
        rm_hdr = t_rm.rows[0].cells
        rm_widths = [Inches(1.8), Inches(1.5), Inches(3.7)]
        for i, h_text in enumerate(["Owner", "Responsibility Area", "Execution Scope & Technical Deliverables"]):
            rm_hdr[i].width = rm_widths[i]
            set_cell_background(rm_hdr[i], "1F4E79" if "Jay" in d_owner else "2E7D32")
            set_cell_margins(rm_hdr[i])
            p = rm_hdr[i].paragraphs[0]
            run = p.add_run(h_text)
            run.bold = True
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
        r1 = t_rm.add_row().cells
        for i, val in enumerate([d_owner, "Development", d_desc]):
            r1[i].width = rm_widths[i]
            set_cell_background(r1[i], "F0F4F8")
            set_cell_margins(r1[i])
            p = r1[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
                
        r2 = t_rm.add_row().cells
        for i, val in enumerate([q_owner, "Testing & QA Audit", q_desc]):
            r2[i].width = rm_widths[i]
            set_cell_background(r2[i], "E8F5E9")
            set_cell_margins(r2[i])
            p = r2[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32)
                
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # SECTION 5: TECHNICAL PLATFORM & INFRASTRUCTURE OWNERSHIP
    # -------------------------------------------------------------
    add_sec_header(doc, "5. Technical Infrastructure, Stack & Production Ownership", level=1)
    
    add_justified_para(
        doc,
        "The entire platform across Product 1, Product 2, and Product 3 is built upon a modern, high-performance web architecture: Next.js + TypeScript for Frontend, NestJS + Express for Backend APIs, PostgreSQL + Prisma / Supabase for Database, and Vercel / Railway for Deployment Infrastructure.",
        bold_prefix="Core Architecture Stack: "
    )
    
    t_inf_master = doc.add_table(rows=1, cols=3)
    t_inf_master.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_inf_master.autofit = False
    set_table_borders(t_inf_master)
    
    im_hdr = t_inf_master.rows[0].cells
    im_widths = [Inches(2.0), Inches(2.5), Inches(2.5)]
    for i, title in enumerate(["Technical Layer", "Jay Kumar (Development Ownership)", "Bharathi (QA Audit Ownership)"]):
        im_hdr[i].width = im_widths[i]
        set_cell_background(im_hdr[i], "1F4E79")
        set_cell_margins(im_hdr[i])
        p = im_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    infra_master_rows = [
        ("Frontend Architecture (Next.js / TypeScript)", "State management, Routing, Shared UI Components, Client-side validation, API client integration.", "UI responsiveness, Cross-browser compatibility, Form error rendering, Navigation smoothness."),
        ("Backend Architecture (NestJS / Express)", "Controller setup, Service layer, DTO validation pipelines, JWT Auth, Role Guard middleware, Business logic.", "API functional testing, Status code verification, Token expiration testing, Access boundary checks."),
        ("Database Layer (PostgreSQL / Prisma / Supabase)", "Schema design, Data relations, Indexing, Constraints, Migrations, Query performance optimization.", "Data integrity validation, Foreign key constraint testing, Duplicate record prevention testing."),
        ("DevOps & Deployment (Vercel / Railway)", "Environment configuration, CI/CD pipeline, Docker containerization, Build verification, Monitoring.", "Post-deployment smoke testing, Sanity testing of production builds, Live environment workflow checks.")
    ]
    for layer, dev_scope, qa_scope in infra_master_rows:
        r_cells = t_inf_master.add_row().cells
        for i, val in enumerate([layer, dev_scope, qa_scope]):
            r_cells[i].width = im_widths[i]
            set_cell_background(r_cells[i], "F0F4F8" if i != 2 else "E8F5E9")
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # SECTION 6: CONTROLLED DELIVERY & BUG MANAGEMENT CYCLE
    # -------------------------------------------------------------
    add_sec_header(doc, "6. Controlled Development → QA → Fix → Retest Delivery Cycle", level=1)
    
    add_justified_para(
        doc,
        "To guarantee zero critical production bugs across Product 1, Product 2, and Product 3, every feature follows a formal 6-stage quality gate before final release.",
        bold_prefix="Quality Gate Policy: "
    )
    
    t_deliv = doc.add_table(rows=1, cols=3)
    t_deliv.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_deliv.autofit = False
    set_table_borders(t_deliv)
    
    d_hdr = t_deliv.rows[0].cells
    d_widths = [Inches(1.8), Inches(1.8), Inches(3.4)]
    for i, title in enumerate(["Delivery Stage", "Responsible Owner", "Execution Standard & Deliverable"]):
        d_hdr[i].width = d_widths[i]
        set_cell_background(d_hdr[i], "1F4E79")
        set_cell_margins(d_hdr[i])
        p = d_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    deliv_steps = [
        ("1. Technical Planning & UI/DB", "Jay Kumar", "Designing database schemas, creating API endpoints, NestJS services, and Next.js frontend pages."),
        ("2. Developer Self-Verification", "Jay Kumar", "Testing happy paths, API payloads, DTO validations, database constraints, and basic UI flow."),
        ("3. QA Sign-Off Request", "Jay Kumar", "Notifying QA that the feature is compiled, deployed to staging, and ready for functional testing."),
        ("4. QA & Functional Testing", "Bharathi", "Testing role permissions, UI responsiveness, edge cases, invalid inputs, and workflow integrity."),
        ("5. Bug Logging & Reporting", "Bharathi", "Filing detailed bug reports with Module, Steps to Reproduce, Expected vs Actual, Screenshots, & Severity."),
        ("6. Root Cause Fix & Retest", "Jay Kumar & Bharathi", "Jay fixes the root cause and deploys update. Bharathi re-tests and conducts full regression before release.")
    ]
    for stage, owner, spec in deliv_steps:
        r_cells = t_deliv.add_row().cells
        for i, val in enumerate([stage, owner, spec]):
            r_cells[i].width = d_widths[i]
            set_cell_background(r_cells[i], "F0F4F8" if "Jay" in owner else "E8F5E9")
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # SECTION 7: FINAL EXECUTIVE SIGN-OFF & DIWALI COMMITMENT
    # -------------------------------------------------------------
    add_sec_header(doc, "7. Final Executive Commitment & Handover Request", level=1)
    
    add_justified_para(
        doc,
        "This master document represents the binding operational blueprint for ISML Portal V3 (Product 1), College LMS Edition (Product 2), and School LMS Edition (Product 3). By enforcing the technical division between Jay Kumar (Lead Development Owner) and Bharathi (Student & QA Testing Owner), the engineering team maintains full control over product velocity and software reliability.",
        bold_prefix="Management Summary: "
    )
    
    make_callout(
        doc,
        "FORMAL DIWALI HOLIDAY HANDOVER GUARANTEE",
        "The engineering team hereby submits this official schedule request: Product 1 (ISML Portal V3) will be delivered on October 4, 2026. "
        "Product 2 (College LMS) and Product 3 (School LMS) will be delivered on October 31, 2026. "
        "Including the 1-Week Contingency Backup Window, ALL three products are guaranteed to be fully completed, tested, and handed over prior to the Diwali Holidays.",
        bg_hex="FFF8E1",
        border_hex="FF8F00"
    )
    
    output_filename = "d:/ISML/resource_agent/ISML_Portal_and_LMS_Grand_Master_Development_Roadmap.docx"
    doc.save(output_filename)
    print(f"Grand Master document successfully created at: {output_filename}")

if __name__ == "__main__":
    build_grand_master_doc()
