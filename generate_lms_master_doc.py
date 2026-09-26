import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_header(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79) # Blue
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32) # Green
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
    elif level == 3:
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(0x6A, 0x1B, 0x9A) # Purple
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
    return p

def add_body_p(doc, text, bold_prefix="", space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.size = Pt(12)
        r_pre.font.name = 'Calibri'
        r_pre.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    r_text = p.add_run(text)
    r_text.font.size = Pt(12)
    r_text.font.name = 'Calibri'
    r_text.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return p

def make_callout_box(doc, text, title="KEY ARCHITECTURAL OVERVIEW", bg_hex="F0F4F8", border_hex="1F4E79"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=160, bottom=160, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}\n")
    run_t.bold = True
    run_t.font.size = Pt(12)
    run_t.font.name = 'Calibri'
    run_t.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    
    run_txt = p.add_run(text)
    run_txt.font.size = Pt(11)
    run_txt.font.name = 'Calibri'
    run_txt.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(6)

def generate_lms_master_doc():
    doc = docx.Document()
    
    # Standard margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Calibri'
    normal_font.size = Pt(12)
    normal_font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("ISML LMS — COMPLETE DEVELOPMENT & RESPONSIBILITY STRUCTURE")
    title_run.bold = True
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    title_p.paragraph_format.space_after = Pt(4)
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Master Software Engineering & QA Responsibility Matrix (Product 2: College LMS & Product 3: School LMS)")
    sub_run.font.size = Pt(13)
    sub_run.italic = True
    sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    sub_p.paragraph_format.space_after = Pt(16)
    
    make_callout_box(
        doc,
        "ISML LMS is engineered as a unified, dual-product Learning Management Platform serving Product 2 (College LMS) and Product 3 (School LMS). "
        "The architecture follows a shared-core model (90% common codebase) with configurable product layers. Development ownership is split between "
        "Jay Kumar (Lead Technical Architect & Full-Stack Developer) and Bharathi (Student Module Developer & Lead QA Auditor).",
        title="EXECUTIVE PRODUCT ARCHITECTURE & RESPONSIBILITY MODEL"
    )
    
    # -------------------------------------------------------------
    # SECTION 1: PRODUCT OVERVIEW & DUAL-EDITION ARCHITECTURE
    # -------------------------------------------------------------
    add_header(doc, "1. Product Overview & Architecture", level=1)
    
    add_body_p(
        doc,
        "The ISML LMS platform is being developed as a common Learning Management ecosystem engineered to support two distinct product editions: Product 2 (College LMS) and Product 3 (School LMS). Rather than building two completely fragmented, siloed platforms, ISML LMS adopts a shared-core + configurable-product architecture. Approximately 90% of the underlying infrastructure—including authentication, database schemas, API controllers, video streaming, finance, scheduling algorithms, and AI engines—is shared across both editions, ensuring maximum code reuse, fast feature deployment, and seamless platform updates.",
        bold_prefix="Dual-Product Strategy: "
    )
    
    add_body_p(
        doc,
        "Product 2 (College LMS) is customized for higher education workflows. It incorporates university academic structures such as Departments, Programs, Semesters, Degree Tracks, Credit Hours, Faculty Assignments, and Course Electives. Product 3 (School LMS) is tailored for K-12 school education, featuring Grade levels, Section management, Subject assignments, Period-based timetabling, Class Teacher controls, and school-specific AI tools (such as Career Guidance and Speaking Bots).",
        bold_prefix="Product Differentiation: "
    )
    
    # Architecture Table
    t_arch = doc.add_table(rows=1, cols=3)
    t_arch.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_arch.autofit = False
    set_table_borders(t_arch)
    
    a_hdr = t_arch.rows[0].cells
    a_widths = [Inches(1.8), Inches(2.3), Inches(2.4)]
    for i, title in enumerate(["Platform Layer", "Product Scope / Edition", "Technical Capabilities & Features"]):
        a_hdr[i].width = a_widths[i]
        set_cell_background(a_hdr[i], "1F4E79")
        set_cell_margins(a_hdr[i])
        p = a_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    arch_rows = [
        ("Shared Core LMS (90%)", "Common Platform Core", "Authentication, Authorization, Prisma Database Schema, NestJS REST APIs, Live Classes (Daily.co/Agora), Video Processing Pipeline, Finance Engine, AI Core, Realtime WebSockets."),
        ("Product 2: College LMS", "Higher Education Edition", "Department Management, Degree Programs, Semester Systems, Credit Allocation, Faculty Assignment, Electives, Semester-wise Attendance & Grading."),
        ("Product 3: School LMS", "K-12 School Edition", "Grades & Sections, Period-based Timetables, Subject Teachers, Class Teachers, Parent Updates, Career Guidance AI, Speaking Practice Bot.")
    ]
    for layer, scope, caps in arch_rows:
        r_cells = t_arch.add_row().cells
        for i, val in enumerate([layer, scope, caps]):
            r_cells[i].width = a_widths[i]
            set_cell_background(r_cells[i], "F0F4F8")
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # SECTION 2: OVERALL TEAM OWNERSHIP & MODEL
    # -------------------------------------------------------------
    add_header(doc, "2. Overall Team Ownership & Responsibility Model", level=1)
    
    add_body_p(
        doc,
        "To guarantee high code quality, zero operational overlap, and rigorous quality assurance before any production deployment, the development responsibility is strictly partitioned between Jay Kumar and Bharathi. Jay Kumar serves as the Lead Development & Technical Owner across the entire platform, while Bharathi serves as the Primary Developer for the Student Module and the Lead QA / Testing Owner for all modules across both Product 2 and Product 3.",
        bold_prefix="Clear Role Partitioning: "
    )
    
    # Ownership Table
    t_own = doc.add_table(rows=1, cols=4)
    t_own.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_own.autofit = False
    set_table_borders(t_own)
    
    o_hdr = t_own.rows[0].cells
    o_widths = [Inches(1.5), Inches(1.8), Inches(1.2), Inches(2.0)]
    for i, title in enumerate(["Team Member", "Primary Technical Ownership", "Theme Color", "Core Operational Responsibilities"]):
        o_hdr[i].width = o_widths[i]
        set_cell_background(o_hdr[i], "1F4E79")
        set_cell_margins(o_hdr[i])
        p = o_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    team_data = [
        ("Jay Kumar", "Full LMS Platform Development Owner", "Blue (#1F4E79)", "Architecture, Database, NestJS Backend, Next.js Frontend, APIs, Super Admin, Manager, Finance, Teacher, Scheduling, Live Room, AI Engines, Deployment, Bug Fixing."),
        ("Bharathi", "Student Module Owner + Lead QA / Testing Owner", "Green (#2E7D32)", "Primary Frontend/UI & Workflow Developer for Student Module. Independent Lead QA Auditor testing every single module, role, and API built by Jay Kumar.")
    ]
    for member, role, color, desc in team_data:
        r_cells = t_own.add_row().cells
        bg = "F0F4F8" if "Jay" in member else "E8F5E9"
        txt_c = RGBColor(0x1F, 0x4E, 0x79) if "Jay" in member else RGBColor(0x2E, 0x7D, 0x32)
        for i, val in enumerate([member, role, color, desc]):
            r_cells[i].width = o_widths[i]
            set_cell_background(r_cells[i], bg)
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = txt_c
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # SECTION 3: STANDARD DEVELOPMENT TO QA BUG CYCLE
    # -------------------------------------------------------------
    add_header(doc, "3. Standard Development → QA → Fix → Release Cycle", level=1)
    
    add_body_p(
        doc,
        "Every single feature delivered in Product 2 (College LMS) and Product 3 (School LMS) must strictly adhere to a controlled development lifecycle. Code is never pushed directly to production by a developer without passing through the formal QA audit pipeline managed by Bharathi.",
        bold_prefix="Controlled Delivery Pipeline: "
    )
    
    # Delivery Cycle Table
    t_cycle = doc.add_table(rows=1, cols=3)
    t_cycle.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_cycle.autofit = False
    set_table_borders(t_cycle)
    
    c_hdr = t_cycle.rows[0].cells
    c_widths = [Inches(1.5), Inches(1.8), Inches(3.2)]
    for i, title in enumerate(["Lifecycle Phase", "Responsible Owner", "Execution Guidelines & Deliverables"]):
        c_hdr[i].width = c_widths[i]
        set_cell_background(c_hdr[i], "1F4E79")
        set_cell_margins(c_hdr[i])
        p = c_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    cycle_steps = [
        ("1. Technical Planning & UI/DB", "Jay Kumar", "System design, Prisma schema creation, REST API controllers, NestJS service logic, and Next.js frontend screen development."),
        ("2. Developer Verification", "Jay Kumar", "Self-testing happy paths, verifying API response payloads, DTO validations, and checking database constraints."),
        ("3. Feature Ready Sign-Off", "Jay Kumar", "Notifying QA that the feature build is compiled, deployed to test server, and ready for functional validation."),
        ("4. QA & Functional Audit", "Bharathi", "Executing comprehensive test cases covering role permissions, UI responsiveness, edge cases, invalid inputs, and workflow integrity."),
        ("5. Bug Logging & Reporting", "Bharathi", "Filing detailed bug tickets containing Module, Steps to Reproduce, Expected vs Actual Behavior, Screenshots, Severity, and Environment."),
        ("6. Root Cause Fix & Retest", "Jay Kumar & Bharathi", "Jay investigates root cause, implements code fix, and deploys update. Bharathi performs re-testing and full regression audit prior to production release.")
    ]
    for phase, owner, details in cycle_steps:
        r_cells = t_cycle.add_row().cells
        for i, val in enumerate([phase, owner, details]):
            r_cells[i].width = c_widths[i]
            set_cell_background(r_cells[i], "F0F4F8" if "Jay" in owner else "E8F5E9")
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # SECTION 4: DETAILED MODULE-BY-MODULE RESPONSIBILITY MATRIX
    # -------------------------------------------------------------
    add_header(doc, "4. Detailed Module-by-Module Development & QA Matrix", level=1)
    
    modules_data = [
        ("4.1 Super Admin Module", 
         "Jay Kumar (Development)", 
         "Platform dashboard, overall LMS statistics, institution creation (School vs College onboarding), multi-tenant database configuration, system-wide role & permission management, global academic settings, system audit logs, and infrastructure health monitoring.",
         "Bharathi (Testing Owner)", 
         "Validating institution onboarding, role privilege boundaries, permission enforcement, academic configuration updates, global course overrides, system audit visibility, error state handling, and security access restrictions."),
         
        ("4.2 Manager Module", 
         "Jay Kumar (Development)", 
         "User account creation (Teachers, Assistant Teachers, Doubt Teachers), User ID generation, academic monitoring dashboards, class & student progress tracking, attendance monitoring, operational approvals, and administrative reports.",
         "Bharathi (Testing Owner)", 
         "Testing user creation workflows, role assignment, class monitoring filters, student attendance reports, operational approval workflows, access boundary checks, and invalid scenario testing."),

        ("4.3 Finance Manager Module", 
         "Jay Kumar (Development)", 
         "Fee structures (Course & Batch fees), student fee assignment, payment status tracking (Paid, Pending, Overdue, Failed), receipt generation, transaction reconciliation, payment gateway integration (Razorpay/Stripe webhooks & signatures), and financial reports.",
         "Bharathi (Testing Owner)", 
         "Testing fee creation, payment link generation, successful & failed payment callbacks, webhook event processing, receipt PDF generation, transaction history accuracy, and financial permission controls."),

        ("4.4 Academic & Resource Manager", 
         "Jay Kumar (Development)", 
         "Core academic hierarchy engine: Language → Course → Level → Module/Unit → Topic → Learning Objective → Skill → Resource. Resource management (Upload, Edit, Approve, Publish, Archive, Metadata indexing, Tagging, and Search APIs).",
         "Bharathi (Testing Owner)", 
         "Validating taxonomy integrity, course level mapping, topic structure, resource upload & approval pipeline, metadata filtering, search index accuracy, and archive state behaviors."),

        ("4.5 Teacher Module", 
         "Jay Kumar (Development)", 
         "Teacher dashboard, assigned courses & batches, class scheduling, live classroom controls (Start/End session, Mute All, Screen Share, Chat), assignment creation, exam management, marks entry, student doubt resolution, and feedback tools.",
         "Bharathi (Testing Owner)", 
         "Validating teacher login, batch filtering, live room control responsiveness, assignment publishing, exam marks submission, doubt response flow, student grading integrity, and recording status visibility."),

        ("4.6 Assistant Teacher Module", 
         "Jay Kumar (Development)", 
         "Assistant Teacher dashboard, assigned class monitoring, student attendance support, assignment review assistance, classroom moderation, resource distribution, and coordination tools.",
         "Bharathi (Testing Owner)", 
         "Testing assistant permission limits (cannot delete courses, cannot alter primary grades), class support workflows, attendance logging verification, and teacher collaboration tools."),

        ("4.7 Doubt Teacher Module", 
         "Jay Kumar (Development)", 
         "Doubt dashboard, real-time doubt queue, subject/topic categorization, student question inspection, text/voice/image response submission, resolution history, and AI-assisted answer suggestions.",
         "Bharathi (Testing Owner)", 
         "Testing doubt submission receipt, queue prioritization, response delivery to student UI, resolution status updates, search & filter functions, and AI answer rendering accuracy."),

        ("4.8 Assistant Doubt Teacher", 
         "Jay Kumar (Development)", 
         "First-level doubt queue monitoring, initial answer submission, resource sharing, escalation workflow (escalate complex doubts to Primary Doubt Teacher), and status tracking.",
         "Bharathi (Testing Owner)", 
         "Testing queue assignment, initial response speed, escalation trigger logic, escalation handoff to Doubt Teacher, and student notification delivery."),

        ("4.9 Student Module (Primary Focus)", 
         "Bharathi (Development Owner)", 
         "Student dashboard UI, My Courses navigation, upcoming classes, live classroom participant view, pre-recorded video player, assignment submission UI, online exam taking UI, doubt raising screen, progress tracking, certificates, and profile management.",
         "Jay Kumar (Backend Integration) & Bharathi (QA Audit)", 
         "Jay provides backend APIs, authentication tokens, WebSocket event feeds, and DB queries. Bharathi builds the entire frontend UI, validates user experience, tests responsive layouts, checks client-side validation, and runs end-to-end regression tests."),

        ("4.10 Common Scheduling Engine", 
         "Jay Kumar (Development)", 
         "Unified scheduling matrix supporting Product 2 (College: Semester → Department → Program → Faculty → Time Slot) and Product 3 (School: Grade → Section → Subject → Teacher → Period). Conflict detection algorithms, rescheduling, and timetable APIs.",
         "Bharathi (Testing Owner)", 
         "Testing schedule creation, time slot conflict detection, teacher overlap prevention, student timetable rendering, class cancellation updates, and Product 2 vs Product 3 workflow differences."),

        ("4.11 Live Classes Engine", 
         "Jay Kumar (Development)", 
         "WebRTC / Daily.co / Agora API integration, secure room token generation, teacher/student authentication, real-time audio/video grid, screen sharing, in-class chat, raise hand feature, participant list, and automated duration logging.",
         "Bharathi (Testing Owner)", 
         "Testing multi-user room entry, token expiration security, camera/mic toggle responsiveness, screen share quality, chat message delivery, network reconnection handling, and session end triggers."),

        ("4.12 Attendance Automation", 
         "Jay Kumar (Development)", 
         "Activity-based automated tracking: Join timestamp capture, leave timestamp capture, disconnect/rejoin tracking, net participation duration calculation, and automatic Present/Absent status updating in Supabase/Prisma.",
         "Bharathi (Testing Owner)", 
         "Testing accuracy of join/leave timestamps, late-join threshold logic, early-leave detection, reconnect duration aggregation, and attendance report correctness."),

        ("4.13 Recording & Pre-Recorded Library", 
         "Jay Kumar (Development)", 
         "Automated recording pipeline: Live Session Completed → Cloud Video Generated → Metadata Extraction → Auto-creation of Pre-Recorded Video item → Video Library publication → HLS Secure Streaming Player with watch history & resume position.",
         "Bharathi (Testing Owner)", 
         "Testing recording generation webhooks, metadata mapping accuracy, video library indexing, secure playback URL authorization, pause/resume position saving, and completion percentage calculation."),

        ("4.14 AI & Intelligent Learning Engine", 
         "Jay Kumar (Development)", 
         "AI Doubt Solver Integration, course-context RAG search, auto question generation, assessment assistance, Speaking Practice Bot (voice evaluation), and Career Guidance AI (Product 3).",
         "Bharathi (Testing Owner)", 
         "Testing AI question input rendering, response latency, prompt accuracy, voice transcription quality in Speaking Bot, feedback rendering, and regression after AI model updates.")
    ]
    
    for title, dev_owner, dev_desc, qa_owner, qa_desc in modules_data:
        add_header(doc, title, level=2)
        
        t_mod = doc.add_table(rows=1, cols=3)
        t_mod.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_mod.autofit = False
        set_table_borders(t_mod)
        
        m_hdr = t_mod.rows[0].cells
        m_widths = [Inches(1.8), Inches(1.5), Inches(3.2)]
        for i, h_text in enumerate(["Role / Owner", "Responsibility Area", "Detailed Execution & Scope"]):
            m_hdr[i].width = m_widths[i]
            set_cell_background(m_hdr[i], "1F4E79" if "Jay" in dev_owner else "2E7D32")
            set_cell_margins(m_hdr[i])
            p = m_hdr[i].paragraphs[0]
            run = p.add_run(h_text)
            run.bold = True
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
        # Development Row
        r1 = t_mod.add_row().cells
        for i, val in enumerate([dev_owner, "Development", dev_desc]):
            r1[i].width = m_widths[i]
            set_cell_background(r1[i], "F0F4F8")
            set_cell_margins(r1[i])
            p = r1[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
                
        # QA Row
        r2 = t_mod.add_row().cells
        for i, val in enumerate([qa_owner, "Testing & QA Audit", qa_desc]):
            r2[i].width = m_widths[i]
            set_cell_background(r2[i], "E8F5E9")
            set_cell_margins(r2[i])
            p = r2[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32)
                
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECTION 5: TECHNICAL PLATFORM & INFRASTRUCTURE OWNERSHIP
    # -------------------------------------------------------------
    add_header(doc, "5. Technical Infrastructure, Stack & Production Ownership", level=1)
    
    add_body_p(
        doc,
        "The underlying technical stack powering both Product 2 (College LMS) and Product 3 (School LMS) requires strict architectural maintenance, automated builds, database optimization, and continuous monitoring.",
        bold_prefix="Core Platform Infrastructure: "
    )
    
    t_infra = doc.add_table(rows=1, cols=3)
    t_infra.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_infra.autofit = False
    set_table_borders(t_infra)
    
    i_hdr = t_infra.rows[0].cells
    i_widths = [Inches(1.8), Inches(1.8), Inches(2.9)]
    for i, title in enumerate(["Infrastructure Layer", "Development Owner (Jay)", "QA Validation Owner (Bharathi)"]):
        i_hdr[i].width = i_widths[i]
        set_cell_background(i_hdr[i], "1F4E79")
        set_cell_margins(i_hdr[i])
        p = i_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    infra_rows = [
        ("Frontend Architecture (Next.js / TypeScript)", "State management, Routing, Layouts, UI Components, Client-side validation, API client integration.", "UI responsiveness, Cross-browser compatibility, Form error state rendering, Navigation smoothness."),
        ("Backend Architecture (NestJS / REST APIs)", "Controller setup, Service layer, DTO validation pipelines, JWT Auth, Role Guard middleware, Business logic.", "API functional testing, Status code verification, Token expiration testing, Access boundary checks."),
        ("Database Layer (PostgreSQL / Prisma / Supabase)", "Schema design, Data relations, Indexing, Constraints, Migrations, Query performance optimization.", "Data integrity validation, Foreign key constraint testing, Duplicate record prevention testing."),
        ("Deployment & DevOps Pipeline", "Environment configuration, Vercel/Railway hosting, Docker containerization, Build verification, Monitoring.", "Post-deployment smoke testing, Sanity testing of production builds, Live environment workflow checks.")
    ]
    for layer, dev_w, qa_w in infra_rows:
        r_cells = t_infra.add_row().cells
        for i, val in enumerate([layer, dev_w, qa_w]):
            r_cells[i].width = i_widths[i]
            set_cell_background(r_cells[i], "F0F4F8" if i != 2 else "E8F5E9")
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # Final Executive Callout
    make_callout_box(
        doc,
        "By enforcing this clear Lead Development Owner (Jay Kumar) and Lead QA / Student Owner (Bharathi) structure across Product 2 (College LMS) and Product 3 (School LMS), ISML LMS achieves maximum engineering velocity, robust code quality, and 100% test coverage prior to production release.",
        title="SUMMARY OF MANAGEMENT EXECUTION & QUALITY COMMITMENT",
        bg_hex="FFF8E1",
        border_hex="FF8F00"
    )
    
    output_filename = "d:/ISML/resource_agent/ISML_LMS_Complete_Development_Responsibility_Structure.docx"
    doc.save(output_filename)
    print(f"Master document successfully created at: {output_filename}")

if __name__ == "__main__":
    generate_lms_master_doc()
