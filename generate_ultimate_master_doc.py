import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_sec_header(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79) # Deep Navy
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
    elif level == 2:
        run.font.size = Pt(13.5)
        run.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32) # Forest Green
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
    elif level == 3:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x6A, 0x1B, 0x9A) # Purple
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
    return p

def add_justified_para(doc, text, bold_prefix="", space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.size = Pt(12)
        r_pre.font.name = 'Calibri'
        r_pre.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        
    r_text = p.add_run(text)
    r_text.font.size = Pt(12)
    r_text.font.name = 'Calibri'
    r_text.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return p

def make_callout(doc, title, text, bg_hex="F0F4F8", border_hex="1F4E79"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=160, bottom=160, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(4)
    
    run_t = p.add_run(f"{title}\n")
    run_t.bold = True
    run_t.font.size = Pt(12)
    run_t.font.name = 'Calibri'
    run_t.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    
    run_txt = p.add_run(text)
    run_txt.font.size = Pt(11)
    run_txt.font.name = 'Calibri'
    run_txt.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(6)

def generate_ultimate_document():
    doc = docx.Document()
    
    # 0.75 in margins for maximum edge-to-edge layout & full content expansion
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Calibri'
    normal_font.size = Pt(12)
    normal_font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    # Title
    t_p = doc.add_paragraph()
    t_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t_run = t_p.add_run("ISML PORTAL & LMS PLATFORM — ULTIMATE MASTER ROADMAP & TEAM RESPONSIBILITY MATRIX")
    t_run.bold = True
    t_run.font.size = Pt(20)
    t_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    t_p.paragraph_format.space_after = Pt(4)
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Exhaustive Technical Specifications for Product 1 (ISML Portal V3), Product 2 (College LMS), & Product 3 (School LMS)")
    sub_run.font.size = Pt(13)
    sub_run.italic = True
    sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    sub_p.paragraph_format.space_after = Pt(16)
    
    make_callout(
        doc,
        "EXECUTIVE TIMELINE & FORMAL RELEASE SCHEDULE",
        "Product 1 (ISML Portal V3) Target Completion Date: October 4, 2026 (Sunday). "
        "Product 2 (College LMS Edition) & Product 3 (School LMS Edition) Target Completion Date: October 31, 2026 (Saturday). "
        "Including a strict 1-Week Contingency Backup Window (November 7), all products are guaranteed to be fully handed over prior to the Diwali Holidays. "
        "Team Ownership is rigorously divided between Jay Kumar (Lead Technical Architect & Full-Stack Development Owner), "
        "Bharathi (Resource Agent Owner & Student Module Developer), and Bharathi + Interns (LSRW Speech & Language Agent Co-Owners).",
        bg_hex="FFF8E1",
        border_hex="FF8F00"
    )

    # -------------------------------------------------------------
    # SECTION 1: MASTER SCHEDULE & RELEASE TIMELINES
    # -------------------------------------------------------------
    add_sec_header(doc, "1. Master Schedule, Milestones & Product Release Timeline", level=1)
    
    add_justified_para(
        doc,
        "The ISML software suite consists of three flagship products engineered to support modern academic institutions and language learners. "
        "Product 1 (ISML Portal V3) focuses on real-time live learning, cloud recording automation, AI resource intelligence, and speech-to-text language evaluation. "
        "Product 2 (College LMS Edition) and Product 3 (School LMS Edition) deliver higher-education and K-12 academic management.",
        bold_prefix="Multi-Product Release Blueprint: "
    )
    
    t_sched = doc.add_table(rows=1, cols=4)
    t_sched.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sched.autofit = False
    set_table_borders(t_sched)
    
    s_hdr = t_sched.rows[0].cells
    s_widths = [Inches(1.8), Inches(1.8), Inches(1.8), Inches(1.6)]
    for i, title in enumerate(["Product Name / Edition", "Target Delivery Date", "Contingency / Backup Buffer", "Diwali Handover Commitment"]):
        s_hdr[i].width = s_widths[i]
        set_cell_background(s_hdr[i], "1F4E79")
        set_cell_margins(s_hdr[i])
        p = s_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    s_data = [
        ("Product 1: ISML Portal V3", "October 4, 2026 (Sunday)", "1-Week Buffer (October 11)", "Guaranteed pre-Diwali Handover"),
        ("Product 2: College LMS Edition", "October 31, 2026 (Saturday)", "1-Week Buffer (November 7)", "Guaranteed pre-Diwali Handover"),
        ("Product 3: School LMS Edition", "October 31, 2026 (Saturday)", "1-Week Buffer (November 7)", "Guaranteed pre-Diwali Handover")
    ]
    for p_name, t_date, c_buf, d_stat in s_data:
        r_cells = t_sched.add_row().cells
        for i, val in enumerate([p_name, t_date, c_buf, d_stat]):
            r_cells[i].width = s_widths[i]
            set_cell_background(r_cells[i], "F0F4F8")
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # SECTION 2: HIGH-LEVEL TEAM OWNERSHIP & WORKFLOW MATRIX
    # -------------------------------------------------------------
    add_sec_header(doc, "2. High-Level Team Ownership & Responsibility Matrix", level=1)
    
    add_justified_para(
        doc,
        "To guarantee high execution speed and absolute quality control, development responsibilities are cleanly partitioned. "
        "Jay Kumar serves as the Lead Technical Architect & Full-Stack Development Owner across the core infrastructure and all administrative modules. "
        "Bharathi serves as the Primary Owner for the Resource Agent and Student Module. "
        "Bharathi + Interns co-own the LSRW Agent (Speech & Language Engine), where Interns actively participate in building exercise datasets, "
        "audio prompt integration, and speech-to-text evaluation workflows.",
        bold_prefix="Core Ownership Model: "
    )
    
    t_team = doc.add_table(rows=1, cols=4)
    t_team.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_team.autofit = False
    set_table_borders(t_team)
    
    tm_hdr = t_team.rows[0].cells
    tm_widths = [Inches(1.5), Inches(1.8), Inches(1.2), Inches(2.5)]
    for i, title in enumerate(["Team Member", "Primary Responsibility Scope", "Theme Color", "Core End-to-End Workflow"]):
        tm_hdr[i].width = tm_widths[i]
        set_cell_background(tm_hdr[i], "1F4E79")
        set_cell_margins(tm_hdr[i])
        p = tm_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    tm_data = [
        ("Jay Kumar", "Live Learning System & Recording Ecosystem", "Blue (#1F4E79)", "Schedule -> Live Room -> Classroom -> Attendance -> Recording -> Processing -> Pre-recorded -> Video Library -> Playback"),
        ("Bharathi", "Resource Intelligence & Resource Agent", "Green (#2E7D32)", "Requirement -> Search -> Discovery -> Metadata -> Classification -> Validation -> Quality -> Duplicate -> Copyright -> Review -> Library"),
        ("Bharathi + Interns", "LSRW Agent (Speech & Language Engine)", "Purple (#6A1B9A)", "Listening -> Speaking -> Reading -> Writing -> AI Evaluation -> Score -> Feedback -> Progress Tracking")
    ]
    for name, scope, color_name, flow in tm_data:
        r_cells = t_team.add_row().cells
        bg_col = "F0F4F8" if "Jay" in name else ("E8F5E9" if "Bharathi" == name else "F3E5F5")
        txt_c = RGBColor(0x1F, 0x4E, 0x79) if "Jay" in name else (RGBColor(0x2E, 0x7D, 0x32) if "Bharathi" == name else RGBColor(0x6A, 0x1B, 0x9A))
        for i, val in enumerate([name, scope, color_name, flow]):
            r_cells[i].width = tm_widths[i]
            set_cell_background(r_cells[i], bg_col)
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = txt_c
                
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # SECTION 3: PRODUCT 1 — ISML PORTAL V3 DEEP DIVE (3 PILLARS)
    # -------------------------------------------------------------
    add_sec_header(doc, "3. Product 1: ISML Portal V3 — Exhaustive Technical Breakdown", level=1)
    
    make_callout(
        doc,
        "V3 CORE PHILOSOPHY: Teach -> Record -> Organize -> Practice -> Evaluate -> Track -> Improve",
        "ISML Portal V3 represents the unified core of Live Learning, AI Resource Intelligence, and Speech/Language Skill Acquisition. "
        "Target Delivery Date: October 4, 2026 (Sunday). Product 1 integrates three primary technical pillars.",
        bg_hex="F0F4F8",
        border_hex="1F4E79"
    )

    # 3.1 Jay Kumar Live System
    add_sec_header(doc, "3.1 Jay Kumar — Live Learning System & Recording Ecosystem", level=2)
    add_justified_para(
        doc,
        "Jay Kumar owns the complete end-to-end engineering of the Live Learning System, cloud recording engine, attendance automation, and pre-recorded video delivery library.",
        bold_prefix="Scope Overview: "
    )
    
    p1_j_table = doc.add_table(rows=1, cols=3)
    p1_j_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    p1_j_table.autofit = False
    set_table_borders(p1_j_table)
    
    pj_hdr = p1_j_table.rows[0].cells
    pj_widths = [Inches(2.0), Inches(1.8), Inches(3.2)]
    for i, title in enumerate(["Module / Sub-system", "Focus Area", "Technical Scope & Detailed Execution Steps"]):
        pj_hdr[i].width = pj_widths[i]
        set_cell_background(pj_hdr[i], "1F4E79")
        set_cell_margins(pj_hdr[i])
        p = pj_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    p1_j_items = [
        ("1. Live Class Scheduling & Session Management", "Class Setup & Lifecycle Control", "Course/Batch/Teacher mapping, Date/Time/Duration configuration, Topic/Objective mapping, Session State transitions (Draft -> Scheduled -> Ready -> Live -> Completed -> Recording Processing -> Recording Available). Handling Failure states (Cancelled, Failed)."),
        ("2. Live Room Creation & Integration", "Room Security & Access Control", "Unique room/session ID generation, Teacher & Student authentication, Secure join URL generation, Session start validation, Reconnection & session end handling."),
        ("3. Teacher Live Classroom", "Teacher Control Room UI", "Start/End session controls, Audio/Video toggles, Screen sharing, Live chat, Participant list, Attendance visibility, Recording status indicator, Connection status monitoring."),
        ("4. Student Live Classroom", "Student Experience & Attendance", "Upcoming class dashboard view, Join validation & waiting room, Live classroom access, Teacher stream viewing, In-class chat, Reconnect & disconnect handling."),
        ("5. Attendance Automation", "Activity-Based Auto Tracking", "Automated join/leave timestamp detection, Participation duration calculation, Rejoin & late join tracking, Auto-calculation of Present/Absent status without manual teacher entry."),
        ("6. Live Session Monitoring & Status", "Real-time Dashboard Visibility", "Management dashboard visibility tracking active sessions, participant counts, teacher connection status, session duration, and recording status in real-time."),
        ("7. Recording Integration", "Cloud Recording & Processing", "Live session auto-recording generation, Availability detection, Recording metadata creation, Processing progress tracking, Failure & retry handling."),
        ("8. Recording -> Pre-recorded Automation", "Zero-Touch Automation Engine", "Automated workflow: Live Class Completed -> Recording Generated -> Metadata Read (Course, Batch, Topic, Teacher) -> Automatic Pre-recorded Session Creation in Video Library."),
        ("9. Pre-recorded Video Library", "Structured Course Library", "Course, batch & topic mapping, Metadata indexing, Secure playback authorization, Student access control, Watch history & resume progress tracking."),
        ("10. Student Video Experience", "Interactive Playback & Tracking", "Video library navigation, Secure player stream, Pause/Resume position saving, Completion percentage tracking, Watch history timeline."),
        ("11. Teacher / Management Dashboard", "Analytics & Administrative Hub", "Comprehensive metrics for Live Classes (Upcoming, Today's, Completed), Attendance statistics, Recording processing state, and Pre-recorded library status."),
        ("12. Testing, Deployment & Production", "QA, Bug Fixes & Deployment", "End-to-end integration testing (Teacher flow, Student flow, Attendance calculation, Recording failures, Playback security), Production deployment & monitoring.")
    ]
    for comp, foc, impl in p1_j_items:
        r_cells = p1_j_table.add_row().cells
        for i, val in enumerate([comp, foc, impl]):
            r_cells[i].width = pj_widths[i]
            set_cell_background(r_cells[i], "F0F4F8")
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 3.2 Bharathi Resource Agent
    add_sec_header(doc, "3.2 Bharathi — Complete Resource Agent & Intelligence Engine", level=2)
    add_justified_para(
        doc,
        "Bharathi has full responsibility for the Resource Agent ecosystem. This module executes automated web discovery, academic taxonomy classification, deduplication, copyright risk analysis, and human review queue orchestration.",
        bold_prefix="Resource Agent Scope: "
    )
    
    p1_b_table = doc.add_table(rows=1, cols=3)
    p1_b_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    p1_b_table.autofit = False
    set_table_borders(p1_b_table)
    
    pb_hdr = p1_b_table.rows[0].cells
    pb_widths = [Inches(2.0), Inches(1.8), Inches(3.2)]
    for i, title in enumerate(["Module / Sub-system", "Focus Area", "Technical Scope & Detailed Execution Steps"]):
        pb_hdr[i].width = pb_widths[i]
        set_cell_background(pb_hdr[i], "2E7D32")
        set_cell_margins(pb_hdr[i])
        p = pb_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    p1_b_items = [
        ("1. Academic Requirement Understanding", "Curriculum & Taxonomy Parser", "Parsing language, course, level, category, skill, topic, learning objective, resource type, difficulty, and academic context."),
        ("2. Resource Discovery", "External Web Search Engine", "Dynamic search query generation, web discovery, candidate filtering, URL collection, and source verification."),
        ("3. Resource Metadata Extraction", "Data Extraction Pipeline", "Automated extraction of title, description, source, provider, language, level, skill, topic, and URL structure."),
        ("4. Academic Classification", "Strict Hierarchy Mapping", "Mapping resources to Language -> Course -> Level -> Category -> Skill -> Topic -> Resource Type structure."),
        ("5. URL & Resource Validation", "Health & Accessibility Check", "Detecting broken links, HTTP redirects, page accessibility, content availability, and metadata completeness."),
        ("6. Duplicate Detection", "Multi-Stage Deduplication", "Deterministic checks: Exact URL -> Canonical URL -> Normalized URL -> URL Hash -> Content Similarity scoring."),
        ("7. Quality & Relevance Analysis", "Evidence-Based Scoring Engine", "Evaluating resource quality, academic relevance, level suitability. Strict Rule: No AI guessing without evidence; flags uncertainty."),
        ("8. Copyright / Usage Analysis", "IP & Compliance Sentinel", "Analyzing license signals, attribution requirements, and usage restrictions. Sends ambiguous legal cases to Human Review."),
        ("9. Human Review Preparation", "Review Queue Orchestration", "Preparing AI analysis results (Classification, Quality, Duplicates, Copyright) for Human Approval / Modification / Rejection."),
        ("10. Resource Library", "Approved Catalog Storage", "Managing indexed resources with multi-faceted filtering (Language, Level, Skill, Topic, Resource Type, Tutor contribution)."),
        ("11. Bulk Resource Import & Reconciliation", "Tutor Data Import & Batch Sync", "Bulk import pipeline, row-by-row duplicate reconciliation, tutor contribution mapping, and failed-row validation error logging.")
    ]
    for comp, foc, impl in p1_b_items:
        r_cells = p1_b_table.add_row().cells
        for i, val in enumerate([comp, foc, impl]):
            r_cells[i].width = pb_widths[i]
            set_cell_background(r_cells[i], "E8F5E9")
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 3.3 Bharathi + Interns LSRW
    add_sec_header(doc, "3.3 Bharathi + Interns — LSRW Agent (Speech & Language AI Engine)", level=2)
    add_justified_para(
        doc,
        "Bharathi and the Interns co-own the LSRW Agent. Bharathi directs the AI evaluation pipeline and architecture, while Interns build practice datasets, curate audio clips, parse transcripts, and test student response workflows.",
        bold_prefix="LSRW Co-Ownership Model: "
    )
    
    p1_l_table = doc.add_table(rows=1, cols=3)
    p1_l_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    p1_l_table.autofit = False
    set_table_borders(p1_l_table)
    
    pl_hdr = p1_l_table.rows[0].cells
    pl_widths = [Inches(2.0), Inches(1.8), Inches(3.2)]
    for i, title in enumerate(["LSRW Skill Area", "Focus Area", "Technical Scope & Detailed Execution Steps"]):
        pl_hdr[i].width = pl_widths[i]
        set_cell_background(pl_hdr[i], "6A1B9A")
        set_cell_margins(pl_hdr[i])
        p = pl_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    p1_l_items = [
        ("🎧 1. Listening Module", "Audio Activity & Question Generation", "Level/Topic-based audio generation, automated question creation, audio playback player, student attempt logging, AI evaluation of comprehension, scoring, and feedback. Interns assist in audio data ingestion and question verification."),
        ("🗣️ 2. Speaking Agent", "Voice Processing & Speech AI Engine", "Audio capture & Speech-to-Text conversion. Multi-dimensional AI Evaluation: Pronunciation, Fluency, Grammar, Vocabulary, Sentence Structure, Response Relevance. Detailed feedback & progress tracking."),
        ("📖 3. Reading Agent", "Passage & Comprehension Evaluation", "Level & topic-based reading content generation, vocabulary-focused passages, comprehension question evaluation, reading score calculation, and improvement recommendations."),
        ("✍️ 4. Writing Agent", "Text Analysis & Essay Scoring", "Writing prompt generation, AI language analysis (Grammar, Vocabulary, Sentence Structure, Coherence, Relevance), instant writing score & suggestion engine."),
        ("🧠 5. Unified LSRW Evaluation & Dashboard", "Academic Progress Ecosystem", "Combining Listening, Speaking, Reading, and Writing scores into a unified AI Feedback engine & Student Academic Dashboard.")
    ]
    for comp, foc, impl in p1_l_items:
        r_cells = p1_l_table.add_row().cells
        for i, val in enumerate([comp, foc, impl]):
            r_cells[i].width = pl_widths[i]
            set_cell_background(r_cells[i], "F3E5F5")
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x6A, 0x1B, 0x9A)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 3.4 Integration Pipelines
    add_sec_header(doc, "3.4 Cross-System Integration Architecture", level=2)
    add_justified_para(
        doc,
        "Product 1 features three automated integration bridges connecting Jay's Live Learning system with Bharathi's Resource Agent and the Bharathi + Interns LSRW Engine.",
        bold_prefix="Automated Pipelines: "
    )
    
    t_int = doc.add_table(rows=1, cols=4)
    t_int.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_int.autofit = False
    set_table_borders(t_int)
    
    in_hdr = t_int.rows[0].cells
    in_widths = [Inches(1.8), Inches(1.5), Inches(1.5), Inches(2.2)]
    for i, title in enumerate(["Integration Pipeline", "Trigger / Input Data", "Responsible System", "Automated Outcome"]):
        in_hdr[i].width = in_widths[i]
        set_cell_background(in_hdr[i], "1F4E79")
        set_cell_margins(in_hdr[i])
        p = in_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    int_items = [
        ("Live Class -> Resource Agent", "Live Class Topic / Learning Objective", "Bharathi (Resource Agent)", "Auto-fetches supporting approved resources from Resource Library for students during live class."),
        ("Live Class -> LSRW Agent", "Live Class Topic / Objective", "Bharathi + Interns (LSRW)", "Generates targeted Listening, Speaking, Reading & Writing practice activities matching live class topics."),
        ("Resource Agent -> LSRW Agent", "Approved Learning Resource", "LSRW AI Evaluation Engine", "Converts approved articles/audio into interactive LSRW exercises with real-time feedback.")
    ]
    for pipe, trig, sys, out in int_items:
        r_cells = t_int.add_row().cells
        for i, val in enumerate([pipe, trig, sys, out]):
            r_cells[i].width = in_widths[i]
            set_cell_background(r_cells[i], "F0F4F8")
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # SECTION 4: PRODUCT 2 (COLLEGE LMS) & PRODUCT 3 (SCHOOL LMS) DEEP DIVE
    # -------------------------------------------------------------
    add_sec_header(doc, "4. Product 2 (College LMS) & Product 3 (School LMS) Complete Specifications", level=1)
    
    add_justified_para(
        doc,
        "Product 2 (College LMS Edition) and Product 3 (School LMS Edition) deliver higher-education and K-12 school management. "
        "The two products share approximately 90% of their core platform architecture while supporting distinct academic structures. Target Completion Date: October 31, 2026.",
        bold_prefix="Shared Core Strategy: "
    )
    
    # 13 LMS Role Modules Expanded
    add_sec_header(doc, "4.1 Exhaustive LMS Roles & System Breakdown", level=2)
    
    full_lms_modules = [
        ("1. Super Admin Module", "Jay Kumar (Development)", "Platform dashboard, overall LMS statistics, institution onboarding (School vs College), multi-tenant DB config, role & permission management, global academic settings, system audit logs, and infrastructure health monitoring.", "Bharathi (QA Testing)", "Validating institution onboarding, role privilege boundaries, permission enforcement, academic updates, global course overrides, system audit visibility, error scenarios, and access restrictions."),
        ("2. Manager Module", "Jay Kumar (Development)", "User account creation (Teachers, Assistant Teachers, Doubt Teachers), User ID generation, teacher management, academic monitoring dashboards, class & student progress tracking, attendance monitoring, and operational approvals.", "Bharathi (QA Testing)", "Testing user creation workflows, role assignment, teacher management, class monitoring filters, student progress reports, operational approval workflows, and access limits."),
        ("3. Finance Manager", "Jay Kumar (Development)", "Fee structure definition, student fee mapping, payment status tracking (Paid, Pending, Failed), receipt PDF generation, transaction reconciliation, payment gateway integration (Razorpay/Stripe webhooks & signatures).", "Bharathi (QA Testing)", "Testing fee creation, payment link generation, successful & failed payment callbacks, webhook event processing, receipt generation, transaction history accuracy, and financial controls."),
        ("4. Academic & Resource Manager", "Jay Kumar (Development)", "Academic hierarchy engine: Language -> Course -> Level -> Module/Unit -> Topic -> Learning Objective -> Skill -> Resource. Resource management (Upload, Edit, Approve, Publish, Archive, Classification, and Search APIs).", "Bharathi (QA Testing)", "Validating taxonomy integrity, course level mapping, topic structure, resource upload & approval pipeline, metadata filtering, search index accuracy, and archive state behaviors."),
        ("5. Teacher Module", "Jay Kumar (Development)", "Teacher dashboard, assigned courses & batches, class scheduling, live classroom controls (Start/End session, Mute All, Screen Share, Chat), assignment creation, exam management, marks entry, doubt resolution, and feedback tools.", "Bharathi (QA Testing)", "Validating teacher login, batch filtering, live room control responsiveness, assignment publishing, exam marks submission, doubt response flow, grading integrity, and recording status."),
        ("6. Assistant Teacher Module", "Jay Kumar (Development)", "Assistant Teacher dashboard, assigned class monitoring, student attendance support, assignment review assistance, classroom moderation, resource distribution, and coordination tools.", "Bharathi (QA Testing)", "Testing assistant permission limits (cannot alter primary grades), class support workflows, attendance logging verification, and teacher collaboration tools."),
        ("7. Doubt Teacher Module", "Jay Kumar (Development)", "Doubt dashboard, real-time doubt queue, subject/topic categorization, student question inspection, text/voice/image response submission, resolution history, and AI-assisted answer suggestions.", "Bharathi (QA Testing)", "Testing doubt submission receipt, queue prioritization, response delivery to student UI, resolution status updates, search & filter functions, and AI answer rendering accuracy."),
        ("8. Assistant Doubt Teacher", "Jay Kumar (Development)", "First-level doubt queue monitoring, initial answer submission, resource sharing, escalation workflow (escalate complex doubts to Primary Doubt Teacher), and status tracking.", "Bharathi (QA Testing)", "Testing queue assignment, initial response speed, escalation trigger logic, escalation handoff to Doubt Teacher, and student notification delivery."),
        ("9. Student Module (Primary Focus)", "Bharathi (Development Owner)", "Student dashboard UI, My Courses navigation, upcoming classes, live classroom participant view, pre-recorded video player, assignment submission UI, online exam taking UI, doubt raising screen, progress tracking, certificates, and profile.", "Jay Kumar (APIs) & Bharathi (QA)", "Jay provides backend APIs, authentication tokens, WebSocket event feeds, and DB queries. Bharathi builds the entire frontend UI, validates user experience, tests responsive layouts, checks client validation, and runs regression tests."),
        ("10. Common Scheduling Engine", "Jay Kumar (Development)", "Unified scheduling matrix supporting Product 2 (College: Semester -> Department -> Program -> Faculty -> Time Slot) and Product 3 (School: Grade -> Section -> Subject -> Teacher -> Period). Conflict detection algorithms, rescheduling, and timetable APIs.", "Bharathi (QA Testing)", "Testing schedule creation, time slot conflict detection, teacher overlap prevention, student timetable rendering, class cancellation updates, and Product 2 vs Product 3 workflow differences."),
        ("11. Live Classes Engine", "Jay Kumar (Development)", "WebRTC / Daily.co / Agora API integration, secure room token generation, teacher/student authentication, real-time audio/video grid, screen sharing, in-class chat, raise hand feature, participant list, and automated duration logging.", "Bharathi (QA Testing)", "Testing multi-user room entry, token expiration security, camera/mic toggle responsiveness, screen share quality, chat message delivery, network reconnection handling, and session end triggers."),
        ("12. Attendance Automation", "Jay Kumar (Development)", "Activity-based automated tracking: Join timestamp capture, leave timestamp capture, disconnect/rejoin tracking, net participation duration calculation, and automatic Present/Absent status updating in Supabase/Prisma.", "Bharathi (QA Testing)", "Testing accuracy of join/leave timestamps, late-join threshold logic, early-leave detection, reconnect duration aggregation, and attendance report correctness."),
        ("13. Recording & Pre-Recorded Library", "Jay Kumar (Development)", "Automated recording pipeline: Live Session Completed -> Cloud Video Generated -> Metadata Extraction -> Auto-creation of Pre-Recorded Video item -> Video Library publication -> HLS Secure Streaming Player with watch history & resume position.", "Bharathi (QA Testing)", "Testing recording generation webhooks, metadata mapping accuracy, video library indexing, secure playback URL authorization, pause/resume position saving, and completion percentage calculation.")
    ]
    
    for r_title, d_owner, d_desc, q_owner, q_desc in full_lms_modules:
        add_sec_header(doc, r_title, level=3)
        t_rm = doc.add_table(rows=1, cols=3)
        t_rm.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_rm.autofit = False
        set_table_borders(t_rm)
        
        rm_hdr = t_rm.rows[0].cells
        rm_widths = [Inches(1.8), Inches(1.5), Inches(3.7)]
        for i, h_text in enumerate(["Owner", "Responsibility Area", "Execution Scope & Detailed Technical Deliverables"]):
            rm_hdr[i].width = rm_widths[i]
            set_cell_background(rm_hdr[i], "1F4E79" if "Jay" in d_owner else "2E7D32")
            set_cell_margins(rm_hdr[i])
            p = rm_hdr[i].paragraphs[0]
            run = p.add_run(h_text)
            run.bold = True
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
        r1 = t_rm.add_row().cells
        for i, val in enumerate([d_owner, "Development", d_desc]):
            r1[i].width = rm_widths[i]
            set_cell_background(r1[i], "F0F4F8")
            set_cell_margins(r1[i])
            p = r1[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
                
        r2 = t_rm.add_row().cells
        for i, val in enumerate([q_owner, "Testing & QA Audit", q_desc]):
            r2[i].width = rm_widths[i]
            set_cell_background(r2[i], "E8F5E9")
            set_cell_margins(r2[i])
            p = r2[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32)
                
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # SECTION 5: TECHNICAL INFRASTRUCTURE & DEVOPS
    # -------------------------------------------------------------
    add_sec_header(doc, "5. Core Platform Infrastructure, Tech Stack & Production Ownership", level=1)
    
    add_justified_para(
        doc,
        "The underlying architecture is standard across Product 1, Product 2, and Product 3: Next.js + TypeScript (Frontend), NestJS + Express (Backend APIs), PostgreSQL + Prisma / Supabase (Database), and Vercel / Railway (DevOps Infrastructure).",
        bold_prefix="Infrastructure Stack: "
    )
    
    t_inf = doc.add_table(rows=1, cols=3)
    t_inf.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_inf.autofit = False
    set_table_borders(t_inf)
    
    i_hdr = t_inf.rows[0].cells
    i_widths = [Inches(2.0), Inches(2.5), Inches(2.5)]
    for i, title in enumerate(["Infrastructure Layer", "Development Owner (Jay Kumar)", "QA Audit Owner (Bharathi)"]):
        i_hdr[i].width = i_widths[i]
        set_cell_background(i_hdr[i], "1F4E79")
        set_cell_margins(i_hdr[i])
        p = i_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    infra_rows = [
        ("Frontend Architecture (Next.js / TypeScript)", "State management, Routing, Shared UI Components, Client-side validation, API client integration.", "UI responsiveness, Cross-browser compatibility, Form error rendering, Navigation smoothness."),
        ("Backend Architecture (NestJS / Express)", "Controller setup, Service layer, DTO validation pipelines, JWT Auth, Role Guard middleware, Business logic.", "API functional testing, Status code verification, Token expiration testing, Access boundary checks."),
        ("Database Layer (PostgreSQL / Prisma / Supabase)", "Schema design, Data relations, Indexing, Constraints, Migrations, Query performance optimization.", "Data integrity validation, Foreign key constraint testing, Duplicate record prevention testing."),
        ("DevOps & Deployment (Vercel / Railway)", "Environment configuration, CI/CD pipeline, Docker containerization, Build verification, Monitoring.", "Post-deployment smoke testing, Sanity testing of production builds, Live environment workflow checks.")
    ]
    for layer, dev_scope, qa_scope in infra_rows:
        r_cells = t_inf.add_row().cells
        for i, val in enumerate([layer, dev_scope, qa_scope]):
            r_cells[i].width = i_widths[i]
            set_cell_background(r_cells[i], "F0F4F8" if i != 2 else "E8F5E9")
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # SECTION 6: QUALITY CONTROL & BUG MANAGEMENT CYCLE
    # -------------------------------------------------------------
    add_sec_header(doc, "6. Controlled Development -> QA -> Fix -> Retest Delivery Cycle", level=1)
    
    add_justified_para(
        doc,
        "Every feature follows a strict 6-stage quality delivery cycle before release to production.",
        bold_prefix="Quality Policy: "
    )
    
    t_del = doc.add_table(rows=1, cols=3)
    t_del.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_del.autofit = False
    set_table_borders(t_del)
    
    d_hdr = t_del.rows[0].cells
    d_widths = [Inches(1.8), Inches(1.8), Inches(3.4)]
    for i, title in enumerate(["Delivery Stage", "Responsible Owner", "Execution Standard & Deliverable"]):
        d_hdr[i].width = d_widths[i]
        set_cell_background(d_hdr[i], "1F4E79")
        set_cell_margins(d_hdr[i])
        p = d_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    deliv_steps = [
        ("1. Technical Planning & UI/DB", "Jay Kumar", "Designing database schemas, creating API endpoints, NestJS services, and Next.js frontend pages."),
        ("2. Developer Self-Verification", "Jay Kumar", "Testing happy paths, API payloads, DTO validations, database constraints, and basic UI flow."),
        ("3. QA Sign-Off Request", "Jay Kumar", "Notifying QA that the feature is compiled, deployed to staging, and ready for functional testing."),
        ("4. QA & Functional Testing", "Bharathi", "Testing role permissions, UI responsiveness, edge cases, invalid inputs, and workflow integrity."),
        ("5. Bug Logging & Reporting", "Bharathi", "Filing detailed bug reports with Module, Steps to Reproduce, Expected vs Actual, Screenshots, & Severity."),
        ("6. Root Cause Fix & Retest", "Jay Kumar & Bharathi", "Jay fixes the root cause and deploys update. Bharathi re-tests and conducts full regression before release.")
    ]
    for stage, owner, spec in deliv_steps:
        r_cells = t_del.add_row().cells
        for i, val in enumerate([stage, owner, spec]):
            r_cells[i].width = d_widths[i]
            set_cell_background(r_cells[i], "F0F4F8" if "Jay" in owner else "E8F5E9")
            set_cell_margins(r_cells[i])
            p = r_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(val)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # SECTION 7: FINAL EXECUTIVE SIGN-OFF & DIWALI COMMITMENT
    # -------------------------------------------------------------
    add_sec_header(doc, "7. Final Executive Commitment & Handover Request", level=1)
    
    make_callout(
        doc,
        "FORMAL RELEASE SCHEDULE & HANDOVER COMMITMENT",
        "Product 1 (ISML Portal V3) Target Delivery Date: October 4, 2026 (Sunday). "
        "Product 2 (College LMS) & Product 3 (School LMS) Target Delivery Date: October 31, 2026 (Saturday). "
        "Including the 1-Week Contingency Backup Window (November 7), all products are guaranteed to be handed over prior to the Diwali Holidays. "
        "Development Ownership is led by Jay Kumar (Architecture, Full-Stack, Live System, DevOps), "
        "Bharathi (Resource Agent & Student Module), and Bharathi + Interns (LSRW Speech & Language Agent Co-Owners).",
        bg_hex="FFF8E1",
        border_hex="FF8F00"
    )
    
    output_filename = "d:/ISML/resource_agent/ISML_Portal_and_LMS_Ultimate_Master_Roadmap.docx"
    doc.save(output_filename)
    print(f"Ultimate document successfully created at: {output_filename}")

if __name__ == "__main__":
    generate_ultimate_document()
