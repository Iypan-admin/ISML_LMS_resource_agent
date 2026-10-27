import os
import logging
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from typing import List, Dict, Any, Optional

logger = logging.getLogger(__name__)

class PPTGenerator:
    """Service to generate professional, ISML-styled PowerPoint presentations."""

    # ISML Color Palette
    PRIMARY_COLOR = RGBColor(30, 58, 138)    # Deep ISML Blue
    SECONDARY_COLOR = RGBColor(14, 165, 233) # Bright Cyan/Teal
    DARK_TEXT = RGBColor(15, 23, 42)         # Slate 900
    LIGHT_TEXT = RGBColor(255, 255, 255)     # White
    MUTED_TEXT = RGBColor(100, 116, 139)     # Slate 500

    @classmethod
    def create_presentation(
        cls,
        title: str,
        subtitle: str,
        slides_data: List[Dict[str, Any]],
        output_path: str,
    ) -> str:
        """Builds a .pptx file with ISML branded title & content slides."""
        prs = Presentation()
        # Set 16:9 widescreen
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        blank_slide_layout = prs.slide_layouts[6]

        # 1. Title Slide
        title_slide = prs.slides.add_slide(blank_slide_layout)
        
        # Background shape
        bg = title_slide.shapes.add_shape(
            1, Inches(0), Inches(0), Inches(13.333), Inches(7.5) # 1 = MSO_SHAPE.RECTANGLE
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = cls.PRIMARY_COLOR
        bg.line.fill.background()

        # Title text box
        tx_box = title_slide.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.333), Inches(2.5))
        tf = tx_box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.bold = True
        p1.font.size = Pt(44)
        p1.font.color.rgb = cls.LIGHT_TEXT
        p1.alignment = PP_ALIGN.LEFT

        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(22)
        p2.font.color.rgb = cls.SECONDARY_COLOR
        p2.alignment = PP_ALIGN.LEFT

        p3 = tf.add_paragraph()
        p3.text = "ISML Language Learning Platform • Faculty Teaching Kit"
        p3.font.size = Pt(14)
        p3.font.color.rgb = RGBColor(203, 213, 225)
        p3.alignment = PP_ALIGN.LEFT

        # 2. Content Slides
        for idx, slide_item in enumerate(slides_data):
            slide = prs.slides.add_slide(blank_slide_layout)

            # Top header bar
            header_bar = slide.shapes.add_shape(
                1, Inches(0), Inches(0), Inches(13.333), Inches(1.3)
            )
            header_bar.fill.solid()
            header_bar.fill.fore_color.rgb = cls.PRIMARY_COLOR
            header_bar.line.fill.background()

            # Slide Header Title
            h_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.25), Inches(11.5), Inches(0.8))
            h_tf = h_box.text_frame
            h_p = h_tf.paragraphs[0]
            h_p.text = slide_item.get("slide_title", f"Slide {idx + 1}")
            h_p.font.bold = True
            h_p.font.size = Pt(28)
            h_p.font.color.rgb = cls.LIGHT_TEXT

            # Slide Main Body Content
            body_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(11.5), Inches(5.0))
            b_tf = body_box.text_frame
            b_tf.word_wrap = True

            points = slide_item.get("bullet_points", [])
            for b_idx, point in enumerate(points):
                bp = b_tf.paragraphs[0] if b_idx == 0 else b_tf.add_paragraph()
                bp.text = f"• {point}"
                bp.font.size = Pt(20)
                bp.font.color.rgb = cls.DARK_TEXT
                bp.space_after = Pt(14)

            # Teacher Notes / Speaking cue
            teacher_cue = slide_item.get("teacher_cue")
            if teacher_cue:
                cue_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.3), Inches(11.5), Inches(0.8))
                c_tf = cue_box.text_frame
                c_p = c_tf.paragraphs[0]
                c_p.text = f"💡 Teacher Cue: {teacher_cue}"
                c_p.font.size = Pt(13)
                c_p.font.italic = True
                c_p.font.color.rgb = cls.MUTED_TEXT

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        prs.save(output_path)
        return output_path
